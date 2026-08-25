#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Generate 开题报告草稿.docx from 开题报告草稿.txt.

School formatting rules (per user, 2026-08-25):
  - 节标题 (1. / 2. / 参考文献): 黑体 小三 (15pt), 段前 0.5 行, 段后 0.5 行
  - 条标题 (1.1 / 2.1 ...):     黑体 四号 (14pt), 段前 0.5 行, 段后 0.5 行
  - 正文:                       宋体 小四 (12pt), 段前 0, 段后 0, 首行缩进 2 字符
  - 数字/英文:                  Times New Roman
  - 行距 1.5 (≈33 行/页), A4, 无页眉
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).parent
SRC = HERE / "开题报告草稿.txt"
OUT = HERE / "开题报告草稿.docx"

HEI = "黑体"
SONG = "宋体"
TNR = "Times New Roman"


def set_run_font(run, ascii_font=TNR, east_font=SONG, size=12, bold=False):
    run.font.name = ascii_font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), ascii_font)
    rFonts.set(qn("w:hAnsi"), ascii_font)
    rFonts.set(qn("w:eastAsia"), east_font)


def set_par_format(p, line=1.5, before=0.0, after=0.0, align=None,
                   first_line_chars=None, hanging_cm=None):
    pf = p.paragraph_format
    pf.line_spacing = line
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if align is not None:
        pf.alignment = align
    if first_line_chars is not None:
        pPr = p._p.get_or_add_pPr()
        ind = pPr.get_or_add_ind()
        ind.set(qn("w:firstLineChars"), str(first_line_chars))
        ind.set(qn("w:firstLine"), "0")
    if hanging_cm is not None:
        pf.left_indent = Cm(hanging_cm)
        pf.first_line_indent = Cm(-hanging_cm)


def main():
    doc = Document()

    # ── Page: A4, default-ish margins, no header ──
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(3.18)
    sec.right_margin = Cm(3.18)
    sec.header_distance = Cm(0)
    # ensure no header content
    sec.header.is_linked_to_previous = False

    # ── Normal style: 宋体 小四 ──
    normal = doc.styles["Normal"]
    normal.font.name = TNR
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), SONG)

    raw = (SRC).read_text(encoding="utf-8")
    lines = [l.rstrip() for l in raw.split("\n")]

    sep_re = re.compile(r"^[═─]{5,}$")
    body_re_h1 = re.compile(r"^\d+．")
    body_re_h2 = re.compile(r"^\d+\.\d+\s")

    in_toc = False  # False | "pending" (after 目录 heading, skip its sep) | True
    for line in lines:
        s = line.strip()
        if not s:
            continue
        if sep_re.match(s):
            if in_toc == "pending":
                in_toc = True
            elif in_toc is True:
                in_toc = False
            continue
        if s.startswith("> "):
            continue  # authoring notes
        if s.startswith("开题报告（草稿"):
            p = doc.add_paragraph()
            r = p.add_run("开题报告（草稿）")
            set_run_font(r, ascii_font=TNR, east_font=HEI, size=22)
            set_par_format(p, line=1.5, before=12, after=6,
                           align=WD_ALIGN_PARAGRAPH.CENTER)
            continue
        if s.startswith("课题名称："):
            p = doc.add_paragraph()
            r = p.add_run(s)
            set_run_font(r, size=14)
            set_par_format(p, line=1.5, before=0, after=12,
                           align=WD_ALIGN_PARAGRAPH.CENTER)
            continue
        if s == "目 录":
            in_toc = "pending"
            p = doc.add_paragraph()
            r = p.add_run("目　录")
            set_run_font(r, ascii_font=TNR, east_font=HEI, size=15)
            set_par_format(p, before=10, after=10,
                           align=WD_ALIGN_PARAGRAPH.CENTER)
            continue
        if in_toc is True:
            p = doc.add_paragraph()
            r = p.add_run(s.lstrip())
            set_run_font(r, size=12)
            set_par_format(p, line=1.25)
            continue
        if s.startswith("参考文献"):
            p = doc.add_paragraph()
            r = p.add_run(s)
            set_run_font(r, ascii_font=TNR, east_font=HEI, size=15)
            set_par_format(p, before=10, after=10)
            continue
        if body_re_h1.match(s):
            p = doc.add_paragraph()
            r = p.add_run(s)
            set_run_font(r, ascii_font=TNR, east_font=HEI, size=15)
            set_par_format(p, before=10, after=10)  # ≈0.5行
            continue
        if body_re_h2.match(s):
            p = doc.add_paragraph()
            r = p.add_run(s)
            set_run_font(r, ascii_font=TNR, east_font=HEI, size=14)
            set_par_format(p, before=9, after=9)  # ≈0.5行
            continue
        # body paragraph
        p = doc.add_paragraph()
        r = p.add_run(s)
        set_run_font(r, size=12)
        if s.startswith("[") and re.match(r"^\[\d+\]", s):
            # reference entry: hanging indent, no first-line indent
            set_par_format(p, line=1.5, hanging_cm=0.85)
        else:
            set_par_format(p, line=1.5, first_line_chars=200)

    doc.save(OUT)
    print(f"Saved {OUT}")


if __name__ == "__main__":
    main()
